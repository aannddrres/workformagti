"""Export routes: compliance CSV, XLSX/PDF compliance reports, the team-stats
PDF, and the async export-job poll/download pair — Phase 9 of the main.py
monolith split.
"""
import csv
import io
import logging
import os
import time
import uuid
from typing import Optional

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy.orm import Session
from starlette.background import BackgroundTask

import models
import security
from config import settings
from database import SessionLocal, get_db
from db_helpers import get_or_404, log_audit
from routers.stats import compute_compliance

router = APIRouter(tags=["exports"])
logger = logging.getLogger("magti")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

_EXPORT_JOB_TTL = 3600  # seconds a finished export stays in the registry
_EXPORT_DIR = os.path.join(settings.UPLOAD_DIR, "exports")

# Guard against a pathologically large export pinning a worker thread or blowing
# the gunicorn timeout. Compliance exports are bounded (users × required readings)
# and sit far below this ceiling at ~600 users. We ERROR rather than silently
# truncate — a partial compliance report is more dangerous than a clear
# "narrow your scope" message.
_EXPORT_MAX_ROWS = 20000


def _guard_export_size(row_count: int) -> None:
    if row_count > _EXPORT_MAX_ROWS:
        raise HTTPException(
            status_code=413,
            detail=(
                f"ექსპორტი ძალიან დიდია ({row_count} ჩანაწერი, ზღვარი "
                f"{_EXPORT_MAX_ROWS}). დააზუსტეთ ფილტრი და სცადეთ თავიდან."
            ),
        )


# Spreadsheet-formula injection (CWE-1236): a free-text field (user name,
# department, title) starting with =, +, -, @, or a tab/CR is read as a
# formula by Excel/LibreOffice on open. A leading quote forces it back to
# literal text in both CSV and XLSX.
_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@", "\t", "\r")


def _sanitize_cell(value):
    if isinstance(value, str) and value.startswith(_FORMULA_TRIGGER_CHARS):
        return "'" + value
    return value


@router.get("/api/export/readings")
def export_readings(
    current_admin: models.User = Depends(security.get_current_system_admin_user),
    db: Session = Depends(get_db)
):
    """Exports required readings compliance status to a CSV file.

    Restricted strictly to system administrators due to containment of employee personal data.
    Logs the export action.

    Access: Restricted to system administrators (admin) only.

    Args:
        current_admin: The authenticated system administrator User.
        db: SQLAlchemy database session.

    Returns:
        A StreamingResponse delivering the exported CSV data.
    """
    # RBAC: this CSV contains personal data (employee names + read timestamps),
    # so it is restricted to the system administrator, not content admins.
    log_audit(db, admin_id=current_admin.id, action="EXPORT", item_type="readings", item_id=0)
    db.commit()

    # Join ReadStatus with User and RequiredReading to output human-readable data.
    # Eligibility comes from compute_compliance() — the same active/non-management
    # rule the dashboard and summary enforce — instead of an unfiltered dump that
    # previously included admins, managers, and inactive users.
    eligible_ids = [r["user_id"] for r in compute_compliance(db)]
    query = db.query(models.ReadStatus, models.User.name, models.RequiredReading.item_type, models.RequiredReading.item_id).join(
        models.User, models.ReadStatus.user_id == models.User.id
    ).join(
        models.RequiredReading, models.ReadStatus.required_reading_id == models.RequiredReading.id
    ).filter(models.ReadStatus.user_id.in_(eligible_ids)).all()
    _guard_export_size(len(query))

    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow(["User ID", "User Name", "Item Type", "Item ID", "Status", "Read At"])

    for rs, user_name, item_type, item_id in query:
        read_at_str = rs.read_at.strftime("%Y-%m-%d %H:%M:%S") if rs.read_at else "N/A"
        row = [rs.user_id, user_name, item_type, item_id, rs.status, read_at_str]
        writer.writerow([_sanitize_cell(v) for v in row])

    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=readings_export.csv"}
    )


# ── XLSX export for compliance readings ───────────────────────────────────
@router.get("/api/export/readings.xlsx")
def export_readings_xlsx(
    background_tasks: BackgroundTasks,
    current_admin: models.User = Depends(security.get_current_system_admin_user),
    db: Session = Depends(get_db),
):
    """Spec-required Excel export of compliance data — admin-only (personal data).

    Falls back gracefully if openpyxl isn't installed (returns 503) so the rest
    of the app keeps working even on a minimal install.
    """
    try:
        from openpyxl import Workbook  # noqa: F401
        from openpyxl.styles import Font, PatternFill  # noqa: F401
    except ImportError:
        raise HTTPException(
            status_code=503,
            detail="openpyxl ბიბლიოთეკა არ არის დაყენებული. გაუშვით: pip install openpyxl",
        )

    log_audit(db, admin_id=current_admin.id, action="EXPORT_XLSX", item_type="readings", item_id=0)
    db.commit()

    rows = db.query(
        models.ReadStatus, models.User.name, models.User.department,
        models.RequiredReading.item_type, models.RequiredReading.item_id,
        models.RequiredReading.due_date,
    ).join(models.User, models.ReadStatus.user_id == models.User.id) \
     .join(models.RequiredReading, models.ReadStatus.required_reading_id == models.RequiredReading.id) \
     .all()
    _guard_export_size(len(rows))

    headers = ["თანამშრომელი", "დეპარტამენტი", "მასალის ტიპი", "მასალის ID", "სტატუსი", "წაკითხვის თარიღი", "ვადა"]
    table_rows = [[
        user_name,
        dept,
        item_type,
        item_id,
        rs.status,
        rs.read_at.strftime("%Y-%m-%d %H:%M") if rs.read_at else "",
        due_date.strftime("%Y-%m-%d") if due_date else "",
    ] for (rs, user_name, dept, item_type, item_id, due_date) in rows]
    return _enqueue_export(background_tasks, table_rows, headers, "Compliance", "xlsx")


# ── PDF export (compliance + team stats) ──────────────────────────────────
# Spec slide 26: Excel/PDF export. XLSX is above; here's the PDF half.
# Georgian Unicode requires a TTF with Georgian glyphs; DejaVu Sans works and
# is installed via fonts-dejavu-core in the Dockerfile. Locally on Windows we
# fall back to Sylfaen.

_GEORGIAN_FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",          # Debian/Ubuntu (Docker)
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",                   # Fedora/RHEL
    "C:/Windows/Fonts/dejavusans.ttf",                          # Windows (if installed)
    "C:/Windows/Fonts/sylfaen.ttf",                             # Windows native Georgian
    os.path.join(BASE_DIR, "static", "fonts", "DejaVuSans.ttf"),
)
_pdf_font_registered: Optional[str] = None


def _register_pdf_font() -> Optional[str]:
    """Register a Unicode-capable TTF for ReportLab; cache the result.

    Returns the registered font name, or None if no font was found.
    """
    global _pdf_font_registered
    if _pdf_font_registered is not None:
        return _pdf_font_registered
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        return None
    for path in _GEORGIAN_FONT_CANDIDATES:
        if os.path.isfile(path):
            try:
                pdfmetrics.registerFont(TTFont("Unicode", path))
                _pdf_font_registered = "Unicode"
                return _pdf_font_registered
            except Exception:
                continue
    return None


def _build_table_pdf(title: str, headers: list[str], rows: list[list[str]]) -> bytes:
    """Render a simple paginated table PDF; raises if reportlab/font is missing."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    font_name = _register_pdf_font()
    if font_name is None:
        raise HTTPException(
            status_code=503,
            detail="PDF Unicode font not found. Install fonts-dejavu-core "
                   "(Debian/Ubuntu) or place DejaVuSans.ttf in static/fonts/.",
        )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Title"], fontName=font_name, fontSize=14)
    body = [Paragraph(title, title_style), Spacer(1, 12)]

    data = [headers] + rows
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#CC0000")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
    ]))
    body.append(table)
    doc.build(body)
    buf.seek(0)
    return buf.getvalue()


@router.get("/api/export/readings.pdf")
def export_readings_pdf(
    background_tasks: BackgroundTasks,
    current_admin: models.User = Depends(security.require_permission(security.PERM_REPORTS_EXPORT)),
    db: Session = Depends(get_db),
):
    """PDF export of compliance readings — admin/manager only via reports.export permission."""
    try:
        import reportlab  # noqa: F401
    except ImportError:
        raise HTTPException(
            status_code=503,
            detail="reportlab ბიბლიოთეკა არ არის დაყენებული. გაუშვით: pip install reportlab",
        )

    log_audit(db, admin_id=current_admin.id, action="EXPORT_PDF", item_type="readings", item_id=0)
    db.commit()

    rows_q = db.query(
        models.ReadStatus, models.User.name, models.User.department,
        models.RequiredReading.item_type, models.RequiredReading.item_id,
        models.RequiredReading.due_date,
    ).join(models.User, models.ReadStatus.user_id == models.User.id) \
     .join(models.RequiredReading, models.ReadStatus.required_reading_id == models.RequiredReading.id) \
     .all()
    _guard_export_size(len(rows_q))

    headers = ["თანამშრომელი", "დეპარტამენტი", "ტიპი", "ID", "სტატუსი", "წაკითხვა", "ვადა"]
    table_rows = [[
        user_name,
        dept or "",
        item_type or "",
        str(item_id),
        rs.status,
        rs.read_at.strftime("%Y-%m-%d %H:%M") if rs.read_at else "",
        due_date.strftime("%Y-%m-%d") if due_date else "",
    ] for (rs, user_name, dept, item_type, item_id, due_date) in rows_q]

    return _enqueue_export(background_tasks, table_rows, headers, "სავალდებულოდ გასაცნობი სტატუსი", "pdf")


@router.get("/api/export/team-stats.pdf")
def export_team_stats_pdf(
    background_tasks: BackgroundTasks,
    current_admin: models.User = Depends(security.require_permission(security.PERM_REPORTS_EXPORT)),
    db: Session = Depends(get_db),
):
    """PDF export of team reading-completion stats by department."""
    try:
        import reportlab  # noqa: F401
    except ImportError:
        raise HTTPException(status_code=503, detail="reportlab არ არის დაყენებული.")

    log_audit(db, admin_id=current_admin.id, action="EXPORT_PDF", item_type="team_stats", item_id=0)
    db.commit()

    # Aggregate by department through compute_compliance() — the same shared
    # formula as the dashboard/summary — instead of a bespoke ReadStatus join
    # that skipped the is_active filter every other compliance view enforces.
    by_dept: dict[str, dict[str, int]] = {}
    for r in compute_compliance(db):
        bucket = by_dept.setdefault(r["department"] or "—", {"total": 0, "read": 0})
        bucket["total"] += r["required_count"]
        bucket["read"] += r["read_count"]

    headers = ["დეპარტამენტი", "სულ მიკუთვნებული", "წაკითხული", "%"]
    table_rows = []
    for dept, vals in sorted(by_dept.items()):
        pct = round(100.0 * vals["read"] / vals["total"], 1) if vals["total"] else 0.0
        table_rows.append([dept, str(vals["total"]), str(vals["read"]), f"{pct}%"])

    return _enqueue_export(background_tasks, table_rows, headers, "გუნდის სტატისტიკა — წაკითხვის პროცენტი", "pdf")


# ── Async file exports (Shape 2a): compile off-request from primitive rows ─
def _build_table_xlsx(title: str, headers: list[str], rows: list[list]) -> bytes:
    """Render a styled single-sheet workbook from primitive rows; raises if
    openpyxl is missing."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Export")[:31]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="CC0000")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
    for i, row in enumerate(rows, start=2):
        for col, val in enumerate(row, start=1):
            ws.cell(row=i, column=col, value=_sanitize_cell(val))
    for col_cells in ws.columns:
        width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(width + 2, 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _async_file_worker(
    job_id: str, table_rows: list, headers: list, title: str, export_type: str
) -> None:
    """Compile an export file from PRIMITIVE rows only (no DB session passed
    in — it opens its own) and record the outcome in export_jobs. Runs in a
    BackgroundTasks worker thread, after the original request's session has
    already closed."""
    db = SessionLocal()
    try:
        if export_type == "xlsx":
            data, ext = _build_table_xlsx(title, headers, table_rows), "xlsx"
        else:
            data, ext = _build_table_pdf(title, headers, table_rows), "pdf"
        os.makedirs(_EXPORT_DIR, exist_ok=True)
        path = os.path.join(_EXPORT_DIR, "export_%s.%s" % (job_id, ext))
        with open(path, "wb") as fh:
            fh.write(data)
        job = db.query(models.ExportJob).filter(models.ExportJob.id == job_id).first()
        if job:
            job.status = "completed"
            job.path = path
            job.expires_at = time.time() + _EXPORT_JOB_TTL
            db.commit()
    except Exception as e:
        logger.warning("export worker failed (job %s): %s", job_id, e)
        job = db.query(models.ExportJob).filter(models.ExportJob.id == job_id).first()
        if job:
            job.status = "failed"
            job.path = None
            job.expires_at = time.time() + _EXPORT_JOB_TTL
            db.commit()
    finally:
        db.close()


def _enqueue_export(
    background_tasks: BackgroundTasks,
    table_rows: list,
    headers: list,
    title: str,
    export_type: str,
) -> dict:
    """Register a job, schedule the off-request build, and return its id."""
    job_id = str(uuid.uuid4())
    db = SessionLocal()
    try:
        db.add(models.ExportJob(
            id=job_id, status="processing", path=None,
            expires_at=time.time() + _EXPORT_JOB_TTL,
        ))
        db.commit()
    finally:
        db.close()
    background_tasks.add_task(
        _async_file_worker, job_id, table_rows, headers, title, export_type
    )
    return {"job_id": job_id}


def _cleanup_export(job_id: str, path: str) -> None:
    """Delete the served file and drop its registry row (runs post-download)."""
    try:
        if path and os.path.exists(path):
            os.unlink(path)
    except OSError:
        pass
    db = SessionLocal()
    try:
        db.query(models.ExportJob).filter(models.ExportJob.id == job_id).delete()
        db.commit()
    finally:
        db.close()


@router.get("/api/export/status/{job_id}")
def get_export_status(
    job_id: str,
    current_admin: models.User = Depends(security.require_permission(security.PERM_REPORTS_EXPORT)),
    db: Session = Depends(get_db),
):
    """Report the state of a background export job."""
    job = get_or_404(db, models.ExportJob, job_id, "საექსპორტო დავალება ვერ მოიძებნა")
    return {"job_id": job_id, "status": job.status}


@router.get("/api/export/download/{job_id}")
def download_export(
    job_id: str,
    current_admin: models.User = Depends(security.require_permission(security.PERM_REPORTS_EXPORT)),
    db: Session = Depends(get_db),
):
    """Serve a completed export, then delete it once the transfer finishes."""
    job = db.query(models.ExportJob).filter(models.ExportJob.id == job_id).first()
    if job is None or job.status != "completed" or not job.path:
        raise HTTPException(status_code=404, detail="ექსპორტი ჯერ არ არის მზად")
    path = job.path
    return FileResponse(
        path,
        filename=os.path.basename(path),
        background=BackgroundTask(_cleanup_export, job_id, path),
    )
